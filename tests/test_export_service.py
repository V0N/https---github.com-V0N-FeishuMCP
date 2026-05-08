from app.service.export_service import generate_excel
from app.service.attendance_service import AttendanceResult, FieldData
from app.service.directory_service import UserInfo
from openpyxl import load_workbook
import io


def make_user(user_id="u_001", name="张三", employee_no="10001", open_id="ou_001", dept=["d_001"]):
    return UserInfo(user_id=user_id, open_id=open_id, name=name, employee_no=employee_no, department_ids=dept)


def make_result(user_id="u_001", name="张三", fields=None):
    if fields is None:
        fields = [
            FieldData(code="50102", title="实际出勤天数", value="21"),
            FieldData(code="50105", title="迟到次数", value="2", is_abnormal=True),
        ]
    return AttendanceResult(name=name, user_id=user_id, fields=fields)


def test_generate_excel_returns_bytes():
    result = make_result()
    user = make_user()
    data = generate_excel([result], {result.user_id: user}, "2026年3月")
    assert isinstance(data, bytes)
    assert len(data) > 0


def test_generate_excel_valid_xlsx():
    result = make_result()
    user = make_user()
    data = generate_excel([result], {result.user_id: user}, "2026年3月")
    wb = load_workbook(io.BytesIO(data))
    assert wb is not None


def test_generate_excel_header_row():
    result = make_result()
    user = make_user()
    data = generate_excel([result], {result.user_id: user}, "2026年3月")
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    header_values = [cell.value for cell in ws[1]]
    assert "姓名" in header_values
    assert "工号" in header_values
    assert "时间范围" in header_values


def test_generate_excel_data_row():
    result = make_result()
    user = make_user()
    data = generate_excel([result], {result.user_id: user}, "2026年3月")
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.max_row == 2
    row2_values = [cell.value for cell in ws[2]]
    assert "张三" in row2_values
    assert "21" in row2_values


def test_generate_excel_empty_data():
    data = generate_excel([], {}, "2026年3月")
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.max_row <= 2


def test_generate_excel_multiple_users():
    result1 = make_result(user_id="u_001", name="张三")
    result2 = make_result(user_id="u_002", name="李四")
    user1 = make_user(user_id="u_001", name="张三")
    user2 = make_user(user_id="u_002", name="李四", open_id="ou_002")
    user_map = {"u_001": user1, "u_002": user2}
    data = generate_excel([result1, result2], user_map, "2026年3月")
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.max_row == 3
