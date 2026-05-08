from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from typing import Optional
import io

from app.service.attendance_service import AttendanceResult
from app.service.directory_service import UserInfo


def generate_excel(
    results: list[AttendanceResult],
    user_map: dict[str, UserInfo],
    time_label: str,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "考勤报表"

    field_titles: list[str] = []
    if results:
        field_titles = [fd.title for fd in results[0].fields]

    header = ["序号", "姓名", "工号", "部门", "时间范围"] + field_titles

    header_font = Font(bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="BDD7EE")
    header_alignment = Alignment(horizontal="center", vertical="center")

    ws.append(header)
    for col_idx, _ in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_idx, result in enumerate(results, start=1):
        user = user_map.get(result.user_id)
        employee_no = user.employee_no if user else ""
        dept = user.department_ids[0] if user and user.department_ids else ""

        field_values = [fd.value for fd in result.fields]

        row = [row_idx, result.name, employee_no, dept, time_label] + field_values
        ws.append(row)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
